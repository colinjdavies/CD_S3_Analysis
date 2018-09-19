import boto3, os
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from utils import human_readable_size
from datetime import datetime
from pricing import get_s3_pricing
import time

PRICING: dict = get_s3_pricing()
FILTER = 'backups'

#Using default profile credentails from AWSCLI currently
client = boto3.client('s3', region_name='eu-west-2')
bucketsList = client.list_buckets()
# Get a list of all bucket names from the response
i = 0
print('0 - ALL BUCKETS')
buckets = bucketsList['Buckets']
for bucket in reversed(buckets):
    i += 1
    name = bucket['Name']
    if FILTER in name:
        print(str(i)+' - '+name)
    else:
        buckets.remove(bucket)
        i -= 1
while True:
    try:
        print()
        bucketSelection = int(input("Which bucket would you like details for? "))
        if int(bucketSelection) >= 0 and int(bucketSelection) <= i:
            break
        else:
            raise Exception('')
    except:
        print("Invalid option")

#TODO Paging for buckets with over 1000 objects

def process_bucket(bucketName, counter=1):
    bucketLocation = client.get_bucket_location(Bucket=bucketName)
    truncatedBucketName = (bucketName[:28] + '..') if len(bucketName) > 30 else bucketName
    print()
    print('Getting objects for bucket: ' + bucketName)
    s3objects = client.list_objects_v2(Bucket=bucketName)
    bucketTotal = 0

    if s3objects['KeyCount'] > 0:
        print('Processing '+str(s3objects['KeyCount'])+' objects for bucket: ' + bucketName)
        ws2 = wb.create_sheet(title=truncatedBucketName)
        ws2.append(["Object Key", "Storage Class", "Size", "Size in Bytes", "Last Modified"])
        firstObject = s3objects['Contents'][0]
        lastModifiedDate = datetime(1970, 1, 1, tzinfo=firstObject["LastModified"].tzinfo)
        for s3object in s3objects['Contents']:
            objectLastModified = s3object["LastModified"]
            if objectLastModified > lastModifiedDate:
                lastModifiedDate = objectLastModified
            objectSize = s3object['Size']
            bucketTotal += objectSize
            dataRow = [s3object["Key"], s3object["StorageClass"], human_readable_size(objectSize), objectSize,
                       objectLastModified]
            ws2.append(dataRow)
        buckettableref = 'A1:E' + str(s3objects['KeyCount'] + 1)
        buckettablename = 'table' + str(counter)
        bucketTab = Table(displayName=buckettablename, ref=buckettableref)
        bucketTab.tableStyleInfo = style
        ws2.add_table(bucketTab)
    else:
        lastModifiedDate = ''
    bucketRegion = 'us-east-1' if bucketLocation['LocationConstraint'] is None else bucketLocation['LocationConstraint']
    totalInGb: float = bucketTotal / 1024 / 1024 / 1024
    costPerGbMo: float = PRICING.get(bucketRegion)
    bucketCost = totalInGb * costPerGbMo
    bucketData = [bucketName, bucketRegion, bucket['CreationDate'], lastModifiedDate,
                  s3objects['KeyCount'], human_readable_size(bucketTotal), bucketCost]
    ws1.append(bucketData)


wb = Workbook()
ws1 = wb.active
ws1.title = "Bucket Summary"
ws1.append(["Name", "Region", "Created Date", "Last Updated", "Total Objects", "Total Size", "Est. Cost ($)"])
# Add a default style with striped rows
style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                       showLastColumn=False, showRowStripes=True, showColumnStripes=False)
if bucketSelection == 0:
    x = len(buckets)
    y = 0
    for bucket in buckets:
        y += 1
        process_bucket(bucket['Name'], y)
else:
    x = 1
    bucket = buckets[len(buckets) - int(bucketSelection)]
    process_bucket(bucket['Name'])

tableRef = 'A1:G'+str(x+1)
summaryTab = Table(displayName="Summary", ref=tableRef)
summaryTab.tableStyleInfo = style
ws1.add_table(summaryTab)

if not os.path.exists('output'):
    os.mkdir('output')
outputFilename = 'output/s3report-'+time.strftime("%Y%m%d-%H%M%S")+'.xlsx'
wb.save(filename=outputFilename)
print()
print('Output saved to ./'+outputFilename)
